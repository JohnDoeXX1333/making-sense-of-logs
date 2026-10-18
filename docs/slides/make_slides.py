"""Build the presentation: docs/slides/making-sense-of-logs.pptx, then export a PDF with LibreOffice.

Usage: python3 docs/slides/make_slides.py
Screenshots in docs/slides/img come from the demo recordings (demo/out/*.webm).
"""
import subprocess
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"
OUT = HERE / "making-sense-of-logs.pptx"

BG = RGBColor(0x0B, 0x0C, 0x0E)
SURFACE = RGBColor(0x16, 0x17, 0x1A)
BORDER = RGBColor(0x2A, 0x2C, 0x31)
TEXT = RGBColor(0xF2, 0xF2, 0xF3)
TEXT2 = RGBColor(0xB4, 0xB6, 0xBC)
TEXT3 = RGBColor(0x8A, 0x8F, 0x98)
ACCENT = RGBColor(0x39, 0x87, 0xE5)
ALERT = RGBColor(0xD9, 0x59, 0x26)
GOOD = RGBColor(0x2E, 0xB8, 0x5C)
FONT = "Noto Sans"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 12


def slide(title: str | None, number: int):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    if title:
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.55), Inches(0.5), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        text(s, title, 0.6, 0.7, 12, 0.8, 30, bold=True)
    text(s, f"Making Sense of Logs  ·  {number} / {TOTAL}", 0.6, 7.0, 6, 0.3, 10, color=TEXT3)
    return s


def text(s, content, x, y, w, h, size=16, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """content: a string, or a list of paragraphs; a paragraph is a string or (string, size, bold, color)."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    paras = content if isinstance(content, list) else [content]
    for i, p in enumerate(paras):
        t, sz, b, c = (p, size, bold, color) if isinstance(p, str) else (p + (size, bold, color)[len(p) - 1:])[:4]
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(sz * 0.45)
        run = para.add_run()
        run.text = t
        run.font.size = Pt(sz)
        run.font.bold = b
        run.font.name = FONT
        run.font.color.rgb = c
    return box


def bullets(s, items, x, y, w, h, size=16, color=TEXT2):
    return text(s, [(f"•  {i}", size, False, color) for i in items], x, y, w, h, size)


def card(s, x, y, w, h, fill=SURFACE):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.adjustments[0] = 0.06
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    r.line.color.rgb = BORDER
    r.line.width = Pt(1)
    r.shadow.inherit = False
    return r


def image(s, name, x, y, w, h, crop=None):
    """Fit the (optionally cropped) screenshot inside the box, keeping its aspect ratio."""
    src = IMG / f"{name}.png"
    im = Image.open(src)
    if crop:
        im = im.crop(crop)
        src = IMG / f"_{name}_{'_'.join(map(str, crop))}.png"
        im.save(src)
    iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    pic = s.shapes.add_picture(str(src), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    pic.line.color.rgb = BORDER
    pic.line.width = Pt(1)
    return pic


def arrow(s, x1, y1, x2, y2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = TEXT3
    c.line.width = Pt(1.75)
    c.line._get_or_add_ln().append(
        c.line._get_or_add_ln().makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"})
    )
    return c


# 1. Title
s = slide(None, 1)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.45), Inches(1.1), Inches(0.08))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()
text(s, "Making Sense of Logs", 0.9, 2.65, 11.5, 1.2, 60, bold=True)
text(s, "One investigation, three perspectives", 0.9, 3.85, 11.5, 0.7, 28, color=TEXT2)
text(s, ["SereneDB challenge · National AI Hackathon 2026",
         ("SereneDB 26.10.1 · DeepSeek V4.1 Flash · Loghub BGL, 4,747,963 log lines", 14, False, TEXT3)], 0.9, 5.4, 11.5, 1.0, 18, color=TEXT2)

# 2. Approach
s = slide("The challenge and our approach", 2)
card(s, 0.6, 1.7, 5.6, 4.9)
text(s, "Same logs, three questions", 0.9, 1.95, 5.1, 0.5, 20, bold=True)
for i, (role, q) in enumerate([("Engineer", "What failed, when, and what came right before?"),
                               ("Team lead", "What keeps coming back, and where?"),
                               ("Manager", "What should we prioritise, and what can the logs not tell us?")]):
    text(s, [(role, 16, True, ACCENT), (q, 16, False, TEXT2)], 0.9, 2.7 + i * 1.25, 5.1, 1.1)
text(s, "Our approach", 6.8, 1.7, 6, 0.5, 20, bold=True)
bullets(s, [
    "SereneDB finds and counts the evidence on the full 4.7M-line log.",
    "An AI agent chooses typed, read-only tools; it never writes SQL.",
    "Every result becomes a numbered fact with its source log lines and the exact SQL.",
    "Three role views explain the same facts; switching roles never changes the evidence.",
    "The server checks every finding: no fact id, no finding; numbers not in the evidence are marked inferred.",
    "When the logs cannot answer (cost, customers, owners), the answer says so.",
], 6.8, 2.3, 6, 4.4, 16)

# 3. System diagram
s = slide("System diagram", 3)
boxes = [
    (0.6, 2.0, 2.0, "Question", "from the web UI, streamed back live"),
    (3.1, 2.0, 2.4, "Planner", "DeepSeek V4.1 Flash picks 2 to 6 tool calls"),
    (6.0, 2.0, 2.6, "Typed tools", "searchLogs · getContext · aggregate, zod-validated"),
    (9.1, 2.0, 3.6, "SereneDB 26.10.1", "4.7M rows, inverted index with BM25, SQL aggregation, read-only"),
    (9.1, 4.3, 3.6, "Evidence bundle", "facts F1..Fn, source log lines, SQL and timings per query"),
    (6.0, 4.3, 2.6, "Priority score", "visible formula over database counts"),
    (3.1, 4.3, 2.4, "Three role views", "engineer, team lead, manager, from the same bundle"),
    (0.6, 4.3, 2.0, "Evidence check", "drops uncited claims, flags unsupported numbers"),
]
for x, y, w, title, sub in boxes:
    card(s, x, y, w, 1.75)
    text(s, [(title, 17, True, TEXT), (sub, 12, False, TEXT2)], x + 0.2, y + 0.16, w - 0.4, 1.5)
for x1, y1, x2, y2 in [(2.6, 2.8, 3.1, 2.8), (5.5, 2.8, 6.0, 2.8), (8.6, 2.8, 9.1, 2.8), (10.9, 3.6, 10.9, 4.3),
                       (9.1, 5.1, 8.6, 5.1), (6.0, 5.1, 5.5, 5.1), (3.1, 5.1, 2.6, 5.1)]:
    arrow(s, x1, y1, x2, y2)
text(s, "The language model never sees the whole log and never writes SQL. Facts come only from SereneDB; the model explains them.",
     0.6, 6.25, 12.1, 0.5, 15, color=TEXT3)

# 4. Tool design
s = slide("Tool design", 4)
tools = [
    ("searchLogs", "Full-text search with filters", "query, phrase, time range, label, level, component, midplane, node", "Inverted index, BM25 ranking, ts_phrase; per-label breakdown"),
    ("getContext", "Timeline around one event", "seed line id, seconds before and after, same node or midplane", "Time-window SQL: events per label and level, sample lines, per-second chart"),
    ("aggregate", "Recurrence, concentration, trend", "group by label, component, level, node, midplane, hour, day or week", "GROUP BY with distinct nodes, active days, FATAL counts and true totals"),
]
for i, (name, what, args, db) in enumerate(tools):
    x = 0.6 + i * 4.15
    card(s, x, 1.7, 3.95, 3.2)
    text(s, [(name, 20, True, ACCENT), (what, 15, True, TEXT), ("Inputs: " + args, 12.5, False, TEXT2), ("SereneDB: " + db, 12.5, False, TEXT2)],
         x + 0.25, 1.9, 3.45, 2.9)
text(s, "Guardrails", 0.6, 5.1, 6, 0.4, 17, bold=True)
bullets(s, [
    "Arguments validated with zod and passed as SQL parameters; read-only queries, client-side time budget (SereneDB does not enforce statement_timeout).",
    "At most 6 tool calls in 3 rounds; the server, not the model, turns results into facts.",
    "Grouped counts use the index only to select rows, then group on the table (a GROUP BY directly on the index overcounted in our tests).",
    "If BM25 scoring fails, search falls back to exact matching and the UI says so.",
], 0.6, 5.5, 12.1, 1.5, 13)

# 5. Data and performance
s = slide("Data and performance", 5)
stats = [("4,747,963", "log lines, every line kept"), ("348,460", "alert lines in 41 categories"), ("3.4 s", "bulk load from Parquet"),
         ("1 to 12 ms", "per search or aggregation query")]
for i, (big, small) in enumerate(stats):
    x = 0.6 + i * 3.1
    card(s, x, 1.7, 2.9, 1.6)
    text(s, [(big, 28, True, TEXT), (small, 13, False, TEXT2)], x + 0.25, 1.85, 2.5, 1.4)
bullets(s, [
    "Loghub BGL (Zenodo 8196385, md5 checked). Parser keeps all lines: 34,470 with an empty message, 306 without a node, 18 wrapped lines are flagged, not dropped.",
    "Times shown as printed in the log (US Pacific); the epoch field is UTC.",
    "Database work for one investigation: 45 to 241 ms. The language model takes about 25 to 35 s; streamed, the evidence appears after about 10 s and the first role view after about 17 s.",
    "Answers to repeated questions are cached, so evidence and numbers stay identical between runs.",
    "Measured on one laptop (Ryzen 7 7745HX, SereneDB in Docker), 2026-10-18.",
], 0.6, 3.6, 12.1, 3.2, 15)

# 6. Engineer
s = slide("Findings for the engineer: the July 11 incident", 6)
image(s, "engineer", 6.2, 1.55, 6.6, 5.3, crop=(210, 60, 1100, 1060))
bullets(s, [
    "14:03:23: 16 KERNPOW 'Power deactivated' events on midplane R05-M0.",
    "14:03:33 to 14:03:39: 148 KERNRTSP 'rts panic! - stopping execution' events on 148 nodes.",
    "Spread: R01-M0 80 nodes, R04-M0 34, R05-M1 34.",
    "The agent states the logs show the order of events, not the cause.",
    "Next checks it recommends: R05-M0 power and cooling at 14:03:23; the job running on the panicked nodes.",
], 0.6, 1.7, 5.3, 5.0, 16)

# 7. Team lead
s = slide("Findings for the team lead: what keeps coming back", 7)
bullets(s, [
    "APPSEV: 49,651 alert lines on 1,024 nodes over 67 days, +216% in the last quarter.",
    "KERNRTSP: 3,983 lines on 3,394 nodes, recurring over 124 days, +9%.",
    "KERNSTOR: 63,491 lines on 22,529 nodes, but only on 5 days: bursts, not a steady problem.",
    "KERNDTLB: 152,734 lines, of which 152,182 come from one node on one day (2005-06-11).",
    "Same evidence as the engineer view; when the model writes a number that is not in the evidence (for example '6 seconds'), the claim is marked inferred.",
], 0.6, 1.7, 5.3, 5.0, 16)
image(s, "teamlead", 6.2, 1.55, 6.6, 5.3, crop=(210, 60, 1100, 1060))

# 8. Manager
s = slide("Findings for the manager: priorities from the logs", 8)
rows = [("Alert category", "Score", "Events", "Nodes", "Active days", "Trend"),
        ("APPSEV", "79", "49,651", "1,024", "67", "+216%"), ("KERNRTSP", "79", "3,983", "3,394", "124", "+9%"),
        ("KERNSTOR", "78", "63,491", "22,529", "5", "+100%"), ("KERNDTLB", "72", "152,734", "541", "11", "n/a"),
        ("KERNTERM", "71", "23,338", "9,524", "28", "-14%")]
tbl = s.shapes.add_table(len(rows), len(rows[0]), Inches(0.6), Inches(1.7), Inches(7.2), Inches(3.0)).table
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = SURFACE if r else RGBColor(0x22, 0x24, 0x28)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
        f = p.runs[0].font
        f.size, f.name, f.bold = Pt(14), FONT, r == 0
        f.color.rgb = TEXT if r else TEXT2
for c, width in enumerate([1.9, 0.9, 1.15, 1.05, 1.2, 1.0]):
    tbl.columns[c].width = Inches(width)
text(s, "Score = 100 × (0.30 events + 0.25 nodes + 0.20 active days + 0.15 fatal share + 0.10 trend), each scaled against the largest category. All inputs are database counts; the trend compares the last quarter with the quarter before.",
     0.6, 4.85, 7.2, 1.0, 12.5, color=TEXT3)
text(s, "What the manager hears", 8.3, 1.7, 4.4, 0.4, 17, bold=True)
bullets(s, [
    "Prioritise APPSEV and KERNRTSP first: widespread and rising.",
    "Treat KERNDTLB as a one-node spike, not a broad issue.",
    "The logs cannot show customer impact, cost or downtime; ask operations for those.",
], 8.3, 2.2, 4.4, 3.5, 15)
text(s, "Asked 'did the panics cause data loss, and what did it cost?', the agent answers that the logs cannot tell, instead of inventing a figure.",
     0.6, 6.0, 12.1, 0.8, 14, color=TEXT2)

# 9. Query evidence
s = slide("Query evidence: every claim links to rows and SQL", 9)
image(s, "drawer", 0.6, 1.55, 6.3, 5.3, crop=(990, 0, 1920, 1080))
bullets(s, [
    "Each finding cites fact ids (F1, F2, ...). A fact holds its value, the source log lines and the query that produced it.",
    "Example: the getContext query around the first panic returned 164 rows in 22 ms; the drawer shows the raw KERNPOW lines and the exact SQL.",
    "The trace lists every query the agent chose, with arguments, row counts and timings.",
    "Smoke test (9 checks) verifies: every finding cites an existing fact, every fact's rows are present, repeated questions give identical facts, unanswerable questions produce no cost or customer claim.",
], 7.2, 1.7, 5.5, 5.0, 14.5)

# 10. Beyond the brief
s = slide("For non-experts and for any log file", 10)
image(s, "insights", 0.6, 1.55, 6.0, 3.6, crop=(210, 60, 1710, 1000))
image(s, "upload_preview", 6.9, 1.55, 5.8, 3.6, crop=(384, 24, 1536, 1056))
bullets(s, [
    "Key takeaways and KPI tiles computed by SQL rules, no AI numbers. Drag a period on the chart: compared with the previous equal period, with the categories and nodes that grew and patterns seen for the first time.",
    "Upload any line-based log: the model proposes a parser, the server tests it on 2,000 lines before loading. Six Loghub formats (HDFS, Linux, Apache, OpenSSH, Zookeeper, Spark): 100% of lines and timestamps parsed in 4 to 9 s.",
], 0.6, 5.35, 12.1, 1.6, 13.5)

# 11. Recommendations
s = slide("Recommendations", 11)
recs = [
    ("Engineer", ["Check R05-M0 power and cooling around 2005-07-11 14:03:23.",
                  "Correlate the 148 panicked nodes with the job that was running.",
                  "Inspect node R30-M0-N9-C:J16-U01, source of 152,182 KERNDTLB lines."]),
    ("Team lead", ["Assign owners to APPSEV and KERNRTSP, the rising, widespread categories.",
                   "Review the KERNSTOR bursts of 2005-06-13 and 2005-06-14.",
                   "Track KERNRTSP recurrence (124 active days)."]),
    ("Manager", ["Prioritise APPSEV and KERNRTSP, then KERNSTOR.",
                 "Do not treat KERNDTLB volume as a fleet-wide problem.",
                 "Get impact, cost and downtime from operations; the logs do not contain them."]),
]
for i, (role, items) in enumerate(recs):
    x = 0.6 + i * 4.15
    card(s, x, 1.7, 3.95, 4.6)
    text(s, role, x + 0.25, 1.9, 3.45, 0.5, 19, bold=True, color=ACCENT)
    bullets(s, items, x + 0.25, 2.5, 3.45, 3.7, 14.5)
text(s, "Taken from the agent's three views of the same evidence; every item cites facts in the app.", 0.6, 6.45, 12.1, 0.4, 13, color=TEXT3)

# 12. Limitations
s = slide("Limitations and what we would do next", 12)
bullets(s, [
    "A new question takes about 25 to 35 s, almost all language model time; streaming shows evidence after about 10 s. Example questions are cached.",
    "Two fresh runs of the same question can choose different queries; cached answers are identical.",
    "BGL has no ownership, customer, cost or downtime data, so those are always reported as unknown.",
    "The priority weights and, for uploaded logs, the alert rule (WARNING or worse, or an error keyword) are our choices; syslog files need the year from the user.",
    "Uploaded samples and facts are sent to the model provider; the upload screen states this.",
    "Measured on one laptop with SereneDB 26.10.1 in Docker (amd64 image).",
], 0.6, 1.7, 7.3, 5.0, 15)
card(s, 8.3, 1.7, 4.4, 4.6)
text(s, [("Next", 19, True, ACCENT),
         ("Run several role writers in parallel on separate provider keys to cut waiting time.", 14, False, TEXT2),
         ("Add the OpenStack logs mentioned for this challenge as a second built-in dataset.", 14, False, TEXT2),
         ("Report the index GROUP BY overcount to SereneDB with our reproduction.", 14, False, TEXT2)],
     8.55, 1.9, 3.95, 4.3)

prs.save(OUT)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(OUT)], check=True,
               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
print(OUT, OUT.with_suffix(".pdf"))
